"""Upload immutable personal reference packs without changing the project references."""
import asyncio
import json
from pathlib import Path
import shutil
import subprocess
import tempfile
from uuid import uuid4
import zipfile

from fastapi import HTTPException
from starlette.responses import JSONResponse

DOCUMENT_LIMIT = 10 * 1024 * 1024
PHOTO_LIMIT = 5 * 1024 * 1024
REQUEST_LIMIT = 30 * 1024 * 1024
DOCUMENT_TYPES = {'.pdf', '.docx', '.doc'}


class UploadBodyLimit:
    """Bound multipart bodies before parsing, including chunked requests."""
    def __init__(self, app):
        self.app = app

    async def __call__(self, scope, receive, send):
        if scope['type'] != 'http' or scope['method'] != 'POST' or scope['path'] != '/api/backgrounds':
            return await self.app(scope, receive, send)
        headers = dict(scope.get('headers', []))
        try:
            declared = int(headers.get(b'content-length', b'0'))
        except ValueError:
            return await JSONResponse({'detail': 'The request length is invalid.'}, 400)(scope, receive, send)
        if declared > REQUEST_LIMIT:
            return await JSONResponse({'detail': 'The upload request must not exceed 30 MB.'}, 413)(scope, receive, send)
        # Disk spooling prevents unbounded RAM and lets multipart parsing clean up normally.
        with tempfile.TemporaryFile() as body:
            size = 0
            while True:
                message = await receive()
                if message['type'] == 'http.disconnect':
                    return
                chunk = message.get('body', b'')
                size += len(chunk)
                if size > REQUEST_LIMIT:
                    return await JSONResponse({'detail': 'The upload request must not exceed 30 MB.'}, 413)(scope, receive, send)
                body.write(chunk)
                if not message.get('more_body', False):
                    break
            body.seek(0)
            remaining = size

            async def replay():
                nonlocal remaining
                if remaining:
                    chunk = body.read(min(65536, remaining))
                    remaining -= len(chunk)
                    return {'type': 'http.request', 'body': chunk, 'more_body': remaining > 0}
                # Support downstream response disconnect listeners after replay ends.
                return await receive()

            if not size:
                return await JSONResponse({'detail': 'The upload must not be empty.'}, 422)(scope, receive, send)
            await self.app(scope, replay, send)


class BackgroundStore:
    def __init__(self, root):
        self.root = Path(root)
        self.entries = {}  # Like tasks, IDs are valid for this server session only.

    def reference(self, ident):
        entry = self.entries.get(ident)
        if not entry:
            raise HTTPException(404, 'Your materials are unavailable or the backend restarted. Please upload them again.')
        if not entry['ready_for_generation']:
            raise HTTPException(422, 'Your personal photo is missing. Upload a complete package before generating the three documents.')
        return self.root / ident / 'reference'

    async def create(self, resume, backgrounds, photo):
        if len(backgrounds) > 5:
            raise HTTPException(422, 'Upload at most five supporting files.')
        self.root.mkdir(parents=True, exist_ok=True, mode=0o700)
        ident = uuid4().hex
        staging = self.root / ('.incoming-' + ident)
        staging.mkdir(mode=0o700)
        try:
            originals = staging / 'originals'
            originals.mkdir()
            items = [(resume, 'resume', DOCUMENT_LIMIT)] + [
                (file, f'background_{i + 1}', DOCUMENT_LIMIT) for i, file in enumerate(backgrounds)]
            if photo:
                items.append((photo, 'photo', PHOTO_LIMIT))
            paths = []
            for upload, name, limit in items:
                suffix = Path(upload.filename or '').suffix.lower()
                allowed = {'.jpg', '.jpeg', '.png'} if name == 'photo' else DOCUMENT_TYPES
                if suffix not in allowed:
                    raise HTTPException(415, 'Documents must be PDF, DOCX or DOC; photos must be JPG or PNG.')
                path = originals / (name + suffix)  # Never use client paths as storage paths.
                size = 0
                with path.open('wb') as target:
                    while chunk := await upload.read(65536):
                        size += len(chunk)
                        if size > limit:
                            raise HTTPException(413, 'Each document must be at most 10 MB and the photo at most 5 MB.')
                        target.write(chunk)
                if not size:
                    raise HTTPException(422, 'Empty files cannot be uploaded.')
                paths.append(path)
            worker = asyncio.create_task(asyncio.to_thread(self._normalize, staging, paths, bool(photo)))
            try:
                warnings = await asyncio.shield(worker)
            except asyncio.CancelledError:
                # Wait for bounded conversion to stop before removing its working files.
                await worker
                raise
            response = {'background_id': ident, 'ready_for_generation': bool(photo), 'warnings': warnings}
            (staging / 'manifest.json').write_text(json.dumps(response, ensure_ascii=False, indent=2), encoding='utf-8')
            staging.rename(self.root / ident)
            self.entries[ident] = response
            return response
        finally:
            if staging.exists():
                shutil.rmtree(staging)

    @staticmethod
    def _docx_text(path):
        from docx import Document
        try:
            with zipfile.ZipFile(path) as archive:
                members = archive.infolist()
                if len(members) > 2000 or sum(p.file_size for p in members) > 50 * 1024 * 1024:
                    raise HTTPException(413, 'The Word file is too large after decompression.')
                if any('vbaproject' in p.filename.lower() for p in members):
                    raise HTTPException(415, 'Word documents containing macros are not accepted.')
            doc = Document(path)
            parts = [doc.part]
            for section in doc.sections:
                for item in (section.header, section.footer, section.first_page_header,
                             section.first_page_footer, section.even_page_header, section.even_page_footer):
                    if item.part not in parts:
                        parts.append(item.part)
            ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
            lines = []
            for part in parts:
                element = part.element
                for paragraph in element.findall('.//w:p', ns):
                    line = ''.join(
                        node.text or '' for node in paragraph.findall('.//w:t', ns))
                    if line.strip():
                        lines.append(line)
            return '\n'.join(lines)
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(422, 'The Word file is damaged, encrypted or not a readable DOCX.') from exc

    @classmethod
    def _text(cls, path, converted):
        if path.suffix == '.docx':
            return cls._docx_text(path), path
        if path.suffix == '.doc':
            if path.read_bytes()[:8] != b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1':
                raise HTTPException(422, 'The file content is not a legacy DOC document.')
            from scripts.check_pages import renderer
            destination = converted / path.stem
            destination.mkdir()
            try:
                result = subprocess.run([renderer(),
                    '-env:UserInstallation=' + (destination / 'profile').as_uri(),
                    '--headless', '--convert-to', 'docx:Office Open XML Text',
                    '--outdir', str(destination), str(path)], capture_output=True, timeout=60)
            except (RuntimeError, OSError, subprocess.TimeoutExpired) as exc:
                raise HTTPException(422, 'DOC conversion failed or timed out. Upload a DOCX or a text-based PDF instead.') from exc
            normalized = destination / (path.stem + '.docx')
            if result.returncode or not normalized.is_file():
                raise HTTPException(422, 'The DOC file could not be converted to DOCX.')
            return cls._docx_text(normalized), normalized
        import pymupdf
        try:
            if path.read_bytes()[:5] != b'%PDF-':
                raise ValueError()
            with pymupdf.open(path) as pdf:
                if pdf.needs_pass:
                    raise ValueError()
                if len(pdf) > 100:
                    raise HTTPException(413, 'PDF files must contain at most 100 pages.')
                texts = []
                for page in pdf:
                    text = page.get_text(sort=True)
                    if not text.strip() and page.get_images():
                        raise HTTPException(422, 'The PDF contains scanned pages without extractable text. Upload a version with selectable text.')
                    texts.append(text)
                return '\n'.join(texts), None
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(422, 'The PDF is damaged, encrypted or invalid.') from exc

    @classmethod
    def _normalize(cls, staging, paths, has_photo):
        from docx import Document
        from PIL import Image, ImageOps
        reference = staging / 'reference'
        reference.mkdir()
        converted = staging / '.converted'
        converted.mkdir()
        warnings = []
        notes = []
        for path in paths:
            if path.stem == 'photo':
                try:
                    with Image.open(path) as image:
                        if image.format not in ('JPEG', 'PNG') or image.width * image.height > 25_000_000:
                            raise ValueError()
                        ImageOps.exif_transpose(image).convert('RGB').save(reference / 'photo.jpg', 'JPEG', quality=95)
                except Exception as exc:
                    raise HTTPException(422, 'The photo is invalid or its dimensions are too large.') from exc
                continue
            text, docx = cls._text(path, converted)
            if not text.strip():
                raise HTTPException(422, 'The document has no readable text. Upload a text-based version.')
            if len(text) > 1_000_000:
                raise HTTPException(413, 'The document contains too much text.')
            if path.suffix == '.pdf':
                warnings.append(f'{path.name} Text was extracted. Multi-column reading order may change; the original file is retained for comparison.')
            if path.stem == 'resume':
                if docx:
                    shutil.copyfile(docx, reference / 'core-resume.docx')
                else:
                    doc = Document()
                    for line in text.splitlines():
                        doc.add_paragraph(line)
                    doc.save(reference / 'core-resume.docx')
                # Include extracted header/footer text too, so the workflow can read identity.
                notes.append(f'## Main resume text (source: originals/{path.name})\n\n{text}')
            else:
                notes.append(f'## Supporting material (source: originals/{path.name})\n\n{text}')
        (reference / 'background-notes.md').write_text('\n\n'.join(notes), encoding='utf-8')
        shutil.copytree(staging / 'originals', reference / 'originals')
        shutil.rmtree(converted)
        if not has_photo:
            warnings.append('A personal photo is required to generate the three documents. Upload a complete package including your photo.')
        return warnings

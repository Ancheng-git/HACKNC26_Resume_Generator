"""Result/preview contract tests with isolated synthetic artifacts; no model calls."""
import asyncio
import hashlib
from io import BytesIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import httpx
from docx import Document
from PIL import Image

from backend.app import create_app
from backend.runner import KINDS, Runner, TaskError
from backend.results import failure_summary

ROOT = Path(__file__).resolve().parents[1]


class ResultTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.base = Path(self.tmp.name).resolve()
        self.directory = self.base / 'task'; self.directory.mkdir()
        (self.directory / 'output').mkdir(); (self.directory / 'preview').mkdir()
        self.runner = Runner(root=ROOT, runs=self.base / 'runs')
        self.task = {'id': 'fixture', 'status': 'succeeded', 'stage': 'completed', 'error': None,
                     'directory': self.directory, 'files': {}, 'hashes': {}, 'previews': {},
                     'preview_hashes': {}, 'ats': {'composite_score': 100, 'must_have_pct': 100}, 'warnings': []}
        for kind in KINDS:
            path = self.directory / 'output' / (kind + '.docx')
            doc = Document(); doc.add_paragraph('Synthetic document: ' + kind); doc.save(path)
            image = self.directory / 'preview' / (kind + '.png')
            Image.new('RGB', (20, 30), 'white').save(image)
            self.task['files'][kind] = path
            self.task['hashes'][kind] = hashlib.sha256(path.read_bytes()).hexdigest()
            self.task['previews'][kind] = image
            self.task['preview_hashes'][kind] = hashlib.sha256(image.read_bytes()).hexdigest()
        self.runner.tasks['fixture'] = self.task
        with patch.dict('os.environ', {'RESUME_FRONTEND_ORIGINS': 'http://localhost:5173'}):
            app = create_app(self.runner)
        self.client = httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url='http://127.0.0.1:8000')

    async def asyncTearDown(self):
        await self.client.aclose(); await self.runner.close(); self.tmp.cleanup()

    async def test_all_three_previews_and_original_downloads(self):
        data = (await self.client.get('/api/generations/fixture')).json()
        for kind in KINDS:
            file = data['files'][kind]
            response = await self.client.get(file['preview_url'], headers={'Origin': 'http://localhost:5173'})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.headers['content-type'], 'image/png')
            self.assertEqual(response.headers['cache-control'], 'no-store')
            self.assertEqual(response.headers['access-control-allow-origin'], 'http://localhost:5173')
            self.assertEqual(response.content, self.task['previews'][kind].read_bytes())
            download = await self.client.get(file['url'])
            self.assertEqual(download.status_code, 200)
            self.assertIn(file['name'], download.headers['content-disposition'])
            self.assertEqual(Document(BytesIO(download.content)).paragraphs[0].text, 'Synthetic document: ' + kind)

    async def test_previews_reject_unknown_failed_running_and_invalid_kind(self):
        self.assertEqual((await self.client.get('/api/generations/unknown/previews/ats_resume')).status_code, 404)
        self.assertEqual((await self.client.get('/api/generations/fixture/previews/not_a_kind')).status_code, 422)
        for status in ('queued', 'running', 'failed'):
            self.task['status'] = status
            response = await self.client.get('/api/generations/fixture/previews/ats_resume')
            self.assertEqual(response.status_code, 409)
            self.assertNotIn('files', (await self.client.get('/api/generations/fixture')).json())

    async def test_changed_or_redirected_files_cannot_be_previewed(self):
        kind = 'ats_resume'; image = self.task['previews'][kind]
        original = image.read_bytes(); image.write_bytes(original + b'changed')
        self.assertEqual((await self.client.get('/api/generations/fixture/previews/' + kind)).status_code, 404)
        image.write_bytes(original)
        self.task['files'][kind].write_bytes(b'changed docx')
        self.assertEqual((await self.client.get('/api/generations/fixture/previews/' + kind)).status_code, 404)
        self.task['hashes'][kind] = hashlib.sha256(b'changed docx').hexdigest()
        outside = self.base / 'outside.png'; outside.write_bytes(original)
        image.unlink(); image.symlink_to(outside)
        self.assertEqual((await self.client.get('/api/generations/fixture/previews/' + kind)).status_code, 404)
        self.task['previews'][kind] = outside
        self.assertEqual((await self.client.get('/api/generations/fixture/previews/' + kind)).status_code, 404)

    async def test_ats_failure_returns_grounded_report_feedback_without_logs(self):
        (self.directory / 'report.md').write_text('''# Report
Contact: PRIVATE PERSON private@example.test
## Genuine gaps
- 4+ years commercial experience: not demonstrated in uploaded materials.
- Databricks: no source evidence.
## Files
/Users/private/material.docx
''')
        (self.directory / 'completion.json').write_text(json.dumps({'revision_passes': 0, 'gaps': 'Databricks'}))
        (self.directory / 'claude.stderr.log').write_text('PRIVATE INTERNAL LOG CONTENT')
        self.task['validation_score'] = {'composite_score': 93, 'keyword_match': {'must_have_pct': 83, 'must_have_missing': ['Databricks']}}
        async def fail(*args):
            raise TaskError('ATS coverage below threshold', 'ats')
        with patch.object(self.runner, '_generate', side_effect=fail):
            await self.runner._execute(self.task, 'Fixture', 'Synthetic JD')
        data = (await self.client.get('/api/generations/fixture')).json()
        self.assertEqual(data['status'], 'failed'); self.assertNotIn('files', data)
        failure = data['failure']
        self.assertEqual(failure['checks']['must_have_pct'], 83)
        self.assertEqual(failure['checks']['revision_passes'], 0)
        self.assertTrue(failure['report_available'])
        self.assertIn('Databricks', json.dumps(failure['report_summary']))
        self.assertIn('Databricks', [area['area'] for area in failure['focus_areas']])
        self.assertIn('Role seniority and commercial experience', [area['area'] for area in failure['focus_areas']])
        for private in ('PRIVATE PERSON', 'private@example.test', '/Users/private', 'PRIVATE INTERNAL LOG'):
            self.assertNotIn(private, json.dumps(data))
        self.assertEqual((await self.client.get('/api/generations/fixture/files/ats_resume')).status_code, 409)

    async def test_authentication_and_timeout_are_not_skill_gap_assessments(self):
        async def fail(*args):
            raise TaskError('Login unavailable', 'authentication')
        with patch.object(self.runner, '_generate', side_effect=fail):
            await self.runner._execute(self.task, 'Fixture', 'Synthetic JD')
        failure = self.runner.public('fixture')['failure']
        self.assertFalse(failure['report_available'])
        self.assertIn('claude auth status', failure['focus_areas'][0]['action'])
        self.assertNotIn('practice', json.dumps(failure))
        self.runner.timeout = 0.01
        async def slow(*args):
            await asyncio.sleep(1)
        with patch.object(self.runner, '_generate', side_effect=slow):
            await self.runner._execute(self.task, 'Fixture', 'Synthetic JD')
        self.assertEqual(self.runner.public('fixture')['failure']['category'], 'timeout')

    def test_bad_or_escaping_reports_are_bounded_and_ignored(self):
        outside = self.base / 'private.md'; outside.write_text('PRIVATE CONTACT')
        (self.directory / 'report.md').symlink_to(outside)
        self.assertFalse(failure_summary(self.task)['report_available'])
        (self.directory / 'report.md').unlink()
        (self.directory / 'report.md').write_text('x' * (128 * 1024 + 1))
        self.assertFalse(failure_summary(self.task)['report_available'])
        (self.directory / 'completion.json').write_text('not json')
        self.assertEqual(failure_summary(self.task)['report_summary'], [])
        self.task['validation_score'] = {'composite_score': float('nan'), 'keyword_match': None}
        feedback = failure_summary(self.task)
        self.assertIsNone(feedback['checks']['composite_score'])
        self.assertEqual(feedback['checks']['missing_must_have'], [])


if __name__ == '__main__':
    unittest.main()

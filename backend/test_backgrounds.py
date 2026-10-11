import asyncio
from io import BytesIO
import json
import os
from pathlib import Path
import subprocess
import shutil
import sys
import tempfile
import unittest
from unittest.mock import patch

import httpx
from docx import Document
from PIL import Image
import pymupdf

ROOT=Path(__file__).resolve().parent.parent
sys.path.insert(0,str(ROOT))
from backend.app import create_app
from backend.backgrounds import BackgroundStore, REQUEST_LIMIT
from backend.runner import Runner
from scripts.check_pages import renderer


def word(text='Synthetic Alice | alice@example.test | Experience 2021-2024 | 42%'):
    d=Document(); d.add_paragraph(text)
    d.add_table(rows=1, cols=1).cell(0,0).text='Table evidence: 123 projects'
    d.sections[0].header.paragraphs[0].text='HEADER CONTACT'
    d.sections[0].footer.paragraphs[0].text='FOOTER RECORD'
    b=BytesIO(); d.save(b); return b.getvalue()


def pdf(text='Synthetic Bob | 2022-2025 | 17%'):
    d=pymupdf.open(); p=d.new_page(); p.insert_text((50,60),text)
    return d.tobytes()


def photo():
    b=BytesIO(); Image.new('RGB',(50,60),'blue').save(b,'PNG'); return b.getvalue()


class UploadTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tmp=tempfile.TemporaryDirectory(); self.base=Path(self.tmp.name)
        self.store=BackgroundStore(self.base/'backgrounds')
        self.runner=Runner(root=ROOT,runs=self.base/'runs',backgrounds=self.store)
        self.client=httpx.AsyncClient(transport=httpx.ASGITransport(app=create_app(self.runner)),base_url='http://127.0.0.1:8000')

    async def asyncTearDown(self):
        await self.client.aclose(); await self.runner.close(); self.tmp.cleanup()

    async def upload(self, content=None, name='resume.docx', extras=None, with_photo=True):
        files=[('resume',(name,content if content is not None else word()))]
        files.extend(extras or [])
        if with_photo: files.append(('photo',('photo.png',photo())))
        return await self.client.post('/api/backgrounds',files=files)

    async def test_generation_denied_command_does_not_publish_files(self):
        response = await self.upload()
        pack = response.json()

        async def fake_process(args, directory, label, input_text=None):
            payload = ({'loggedIn': True} if label == 'auth' else
                       {'subtype': 'success', 'is_error': False,
                        'permission_denials': [{'tool_name': 'Bash'}]})
            (directory / (label + '.stdout.log')).write_text(json.dumps(payload))
            return 0

        with patch.object(self.runner, '_process', side_effect=fake_process):
            response = await self.client.post('/api/generations', json={
                'company': 'Test', 'job_description': 'Python analyst',
                'background_id': pack['background_id']})
            self.assertEqual(response.status_code, 202)
            ident = response.json()['id']
            await self.runner.active
        status = (await self.client.get('/api/generations/' + ident)).json()
        self.assertEqual(status['status'], 'failed')
        self.assertIn('rejected by permission rules', status['error'])
        self.assertNotIn('files', status)
        download = await self.client.get('/api/generations/' + ident + '/files/ats_resume')
        self.assertEqual(download.status_code, 409)

    async def test_docx_upload_headers_tables_originals(self):
        r=await self.upload(extras=[('background_files',('more.docx',word('Extra project: 99 customers')))])
        self.assertEqual(r.status_code,201,r.text)
        data=r.json(); self.assertTrue(data['ready_for_generation']); self.assertEqual(data['warnings'],[])
        ref=self.store.reference(data['background_id'])
        notes=(ref/'background-notes.md').read_text()
        for marker in ['Synthetic Alice','42%','HEADER CONTACT','FOOTER RECORD','123 projects','99 customers']:
            self.assertIn(marker,notes)
        self.assertEqual((ref/'core-resume.docx').read_bytes(),(ref/'originals/resume.docx').read_bytes())
        self.assertTrue((ref/'originals/resume.docx').is_file())
        with Image.open(ref/'photo.jpg') as image:
            self.assertEqual(image.format,'JPEG')

    async def test_pdf_normalization(self):
        r=await self.upload(pdf(), 'resume.pdf',extras=[('background_files',('experience.pdf',pdf('Project 2020-2021 | 30 metrics')))])
        self.assertEqual(r.status_code,201,r.text)
        ref=self.store.reference(r.json()['background_id'])
        self.assertIn('Synthetic Bob',BackgroundStore._docx_text(ref/'core-resume.docx'))
        self.assertIn('30 metrics',(ref/'background-notes.md').read_text())
        self.assertEqual(len(r.json()['warnings']),2)

    async def test_real_legacy_doc_conversion(self):
        source=self.base/'legacy.docx'; source.write_bytes(word('Legacy Synthetic Carol 2019-2023 35%'))
        dest=self.base/'legacy-converted'; dest.mkdir()
        result=await asyncio.to_thread(subprocess.run,[renderer(),'-env:UserInstallation='+(self.base/'lo-profile').as_uri(),'--headless','--convert-to','doc:MS Word 97','--outdir',str(dest),str(source)],capture_output=True,timeout=60)
        self.assertEqual(result.returncode,0)
        legacy=dest/'legacy.doc'; self.assertTrue(legacy.is_file())
        r=await self.upload(legacy.read_bytes(),'resume.doc')
        self.assertEqual(r.status_code,201,r.text)
        self.assertIn('Legacy Synthetic Carol',BackgroundStore._docx_text(self.store.reference(r.json()['background_id'])/'core-resume.docx'))

    async def test_missing_photo_is_saved_but_generation_rejected(self):
        r=await self.upload(with_photo=False)
        self.assertEqual(r.status_code,201)
        data=r.json(); self.assertFalse(data['ready_for_generation'])
        r=await self.client.post('/api/generations',json={'company':'test','job_description':'JD','background_id':data['background_id']})
        self.assertEqual(r.status_code,422)
        self.assertFalse(self.runner.tasks)
        self.assertFalse((self.store.root/data['background_id']/'reference/photo.jpg').exists())

    async def test_invalid_ids_and_restart(self):
        for ident,code in [('not-an-id',422),('a'*32,404),('../../reference',422)]:
            r=await self.client.post('/api/generations',json={'company':'test','job_description':'JD','background_id':ident})
            self.assertEqual(r.status_code,code)
        r=await self.upload()
        newstore=BackgroundStore(self.store.root)
        with self.assertRaises(Exception) as error: newstore.reference(r.json()['background_id'])
        self.assertEqual(error.exception.status_code,404)

    async def test_two_packs_and_legacy_are_isolated(self):
        ids=[]
        for who in ['Alice','Bob']:
            r=await self.upload(word('Synthetic '+who)); ids.append(r.json()['background_id'])
        snapshots=[]
        async def capture(task,company,jd):
            directory,*_=self.runner._prepare(task,company,jd)
            snapshots.append((directory/'reference/background-notes.md').read_text())
            self.assertEqual((directory/'reference/core-resume.docx').read_bytes(),(directory/'reference/originals/resume.docx').read_bytes())
            self.assertTrue((directory/'reference/keywords_schema.json').exists())
            self.assertEqual(list((directory/'reference/style-examples').iterdir()),[])
            task.update(status='failed',stage='failed',error='SIMULATED: snapshot checked')
        self.runner._generate=capture
        for ident in ids:
            r=await self.client.post('/api/generations',json={'company':'Test','job_description':'JD','background_id':ident})
            self.assertEqual(r.status_code,202,r.text); await self.runner.active
        self.assertIn('Alice',snapshots[0]); self.assertNotIn('Bob',snapshots[0])
        self.assertIn('Bob',snapshots[1]); self.assertNotIn('Alice',snapshots[1])
        self.assertNotIn('Ancheng',snapshots[0]+snapshots[1])
        legacy_root = self.base/'legacy-root'
        for name in ('scripts', '.claude'):
            shutil.copytree(ROOT/name, legacy_root/name)
        legacy_reference = legacy_root/'reference'
        legacy_reference.mkdir()
        (legacy_reference/'core-resume.docx').write_bytes(word('Synthetic legacy reference'))
        (legacy_reference/'photo.jpg').write_bytes(photo())
        old={'id':'legacy'}
        with patch.object(self.runner, 'root', legacy_root):
            directory,*_=self.runner._prepare(old,'test','JD')
        self.assertEqual((directory/'reference/core-resume.docx').read_bytes(),(legacy_reference/'core-resume.docx').read_bytes())
        self.assertEqual((directory/'reference/photo.jpg').read_bytes(),(legacy_reference/'photo.jpg').read_bytes())

    async def test_invalid_format_empty_corruption_and_cleanup(self):
        for content,name,expected in [(b'bad','x.exe',415),(b'','x.docx',422),(b'bad','x.docx',422),(b'bad','x.pdf',422),(b'bad','x.doc',422),(word(),'x.pdf',422)]:
            r=await self.upload(content,name)
            self.assertEqual(r.status_code,expected,(name,r.text))
        r=await self.upload(extras=[('background_files',('bad.docx',b'broken'))])
        self.assertEqual(r.status_code,422)
        self.assertEqual(list(self.store.root.iterdir()),[])
        r=await self.upload(name='../../outside.docx')
        self.assertEqual(r.status_code,201)
        self.assertFalse((self.base/'outside.docx').exists())
        self.assertTrue((self.store.root/r.json()['background_id']/'originals/resume.docx').exists())

    async def test_scan_encrypted_empty_pdf_and_photo(self):
        scan=pymupdf.open(); p=scan.new_page(); p.insert_image(p.rect,stream=photo())
        encrypted=pymupdf.open(); encrypted.new_page().insert_text((50,60),'secret')
        secured=encrypted.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256,owner_pw='owner',user_pw='password')
        blank=pymupdf.open(); blank.new_page()
        for payload in [scan.tobytes(),secured,blank.tobytes()]:
            r=await self.upload(payload,'resume.pdf'); self.assertEqual(r.status_code,422,r.text)
        r=await self.upload(with_photo=False,extras=[('photo',('photo.png',b'bad'))]); self.assertEqual(r.status_code,422)

    async def test_file_counts_and_request_fields(self):
        for files in [[('photo',('p.png',photo()))],[('resume',('a.docx',word())),('resume',('b.docx',word()))],[('resume',('r.docx',word())),('unexpected',('x.docx',word()))],[('resume',('r.docx',word()))]+[('background_files',(f'{i}.docx',word())) for i in range(6)]]:
            r=await self.client.post('/api/backgrounds',files=files); self.assertEqual(r.status_code,422,r.text)
        r=await self.client.post('/api/backgrounds',json={'resume':'not a file'}); self.assertEqual(r.status_code,415)
        r=await self.client.post('/api/backgrounds',data={'resume':'text'},files={'background_files':('x.docx',word())}); self.assertEqual(r.status_code,422)

    async def test_size_limits_and_chunked_total(self):
        r=await self.upload(b'x'*(10*1024*1024+1)); self.assertEqual(r.status_code,413)
        r=await self.upload(with_photo=False,extras=[('photo',('p.png',b'x'*(5*1024*1024+1)))]); self.assertEqual(r.status_code,413)
        r=await self.client.post('/api/backgrounds',content=b'x',headers={'content-length':str(REQUEST_LIMIT+1),'content-type':'multipart/form-data; boundary=a'}); self.assertEqual(r.status_code,413)
        async def chunks():
            for _ in range(31): yield b'x'*(1024*1024)
        r=await self.client.post('/api/backgrounds',content=chunks(),headers={'content-type':'multipart/form-data; boundary=a'}); self.assertEqual(r.status_code,413)
        self.assertEqual(list(self.store.root.iterdir()),[])

    async def test_converter_failure_and_timeout(self):
        ole=b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1'+b'test'
        with patch('backend.backgrounds.subprocess.run',side_effect=subprocess.TimeoutExpired('soffice',60)):
            r=await self.upload(ole,'resume.doc'); self.assertEqual(r.status_code,422); self.assertIn('timed out',r.text)
        with patch('backend.backgrounds.subprocess.run',return_value=subprocess.CompletedProcess([],1)):
            r=await self.upload(ole,'resume.doc'); self.assertEqual(r.status_code,422)
        self.assertEqual(list(self.store.root.iterdir()),[])

    async def test_upload_to_real_builders_validation_and_download_without_model(self):
        # Fake only the AI content step; run the existing renderers and validators.
        fake = self.base/'fake-claude'
        fake.write_text('#!' + sys.executable + '\n' + r"""
import sys,json,hashlib,subprocess
from pathlib import Path
from datetime import datetime
from zoneinfo import ZoneInfo
from docx import Document
import pymupdf
if sys.argv[1:3] == ['auth','status']:
    print(json.dumps({'loggedIn':True})); sys.exit(0)
sys.stdin.read()
name=Document('reference/core-resume.docx').paragraphs[0].text
sys.path.insert(0,str(Path('scripts').resolve()))
from build_professional import build as resume_build
from build_ats_resume import build as ats_build
from build_cover_letter_professional import build as letter_build
from check_pages import renderer
c={'name':name,'contact':['fixture@example.test'],'education':[{'degree':'Synthetic degree','school':'Fixture University'}],
'core_expertise':['Python'],'headline':'Synthetic Analyst','profile':['Synthetic analyst with Python experience.'],
'experience':[{'company':'Fixture Company','title':'Analyst','dates':'2021-2024','bullets':['Completed 42 projects.']}]}
date=datetime.now(ZoneInfo('America/New_York')).date().isoformat()
names={'ats_resume':f'Resume_Test_{date}.docx','professional_resume':f'Professional_Test_{date}.docx','professional_cover_letter':f'Professional_CoverLetter_Test_{date}.docx'}
Path('professional_Test.json').write_text(json.dumps(c))
ats_build(c,Path('output')/names['ats_resume'])
resume_build(c,Path('output')/names['professional_resume'])
letter_build({**c,'date':date,'salutation':'Dear Team,','paragraphs':['This is a synthetic test letter.']},Path('output')/names['professional_cover_letter'])
Path('keywords_Test.json').write_text(json.dumps({'must_have':['Python'],'nice_to_have':[]}))
reviews={}
for kind,file in names.items():
    path=Path('output')/file
    subprocess.run([renderer(),'-env:UserInstallation='+Path('preview/profile').resolve().as_uri(),'--headless','--convert-to','pdf','--outdir','preview',str(path)],check=True,capture_output=True,timeout=60)
    image='preview/'+kind+'.png'
    with pymupdf.open(Path('preview')/(path.stem+'.pdf')) as pdf: pdf[0].get_pixmap().save(image)
    reviews[kind]={'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'image':image,'notes':'SIMULATED AI attestation for integration test, not human review'}
Path('completion.json').write_text(json.dumps({'visual_review':reviews,'revision_passes':0,'gaps':''}))
Path('report.md').write_text('Synthetic integration fixture; no model called.')
print(json.dumps({'subtype':'success','is_error':False}))
""")
        fake.chmod(0o700); self.runner.claude=str(fake)
        d=Document()
        for line in ['Synthetic Alice','fixture@example.test','Synthetic degree, Fixture University','Python','Fixture Company Analyst 2021-2024','Completed 42 projects.']:
            d.add_paragraph(line)
        b=BytesIO(); d.save(b)
        pack=(await self.upload(b.getvalue())).json()
        response=await self.client.post('/api/generations',json={'company':'Test','job_description':'Python analyst','background_id':pack['background_id']})
        self.assertEqual(response.status_code,202)
        ident=response.json()['id']
        await asyncio.wait_for(asyncio.shield(self.runner.active),60)
        result=(await self.client.get('/api/generations/'+ident)).json()
        self.assertEqual(result['status'],'succeeded',result)
        self.assertEqual(result['ats']['must_have_pct'],100)
        self.assertTrue(result['warnings'])  # Short synthetic evidence warns without inventing filler.
        for file in result['files'].values():
            response=await self.client.get(file['url'])
            self.assertEqual(response.status_code,200)
            doc=Document(BytesIO(response.content))
            text=' '.join(n.text or '' for n in doc.element.findall('.//{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
            self.assertIn('synthetic alice',text.lower())
            self.assertNotIn('ancheng',text.lower())

    async def test_openapi_and_cors(self):
        schema=(await self.client.get('/openapi.json')).json()
        self.assertIn('/api/backgrounds',schema['paths'])
        self.assertIn('background_id',schema['components']['schemas']['GenerationRequest']['properties'])
        r=await self.client.post('/api/backgrounds',files={'resume':('r.docx',word())},headers={'Origin':'https://evil.example'})
        self.assertEqual(r.status_code,403)
        with patch.dict(os.environ, {'RESUME_FRONTEND_ORIGINS':'http://localhost:5173'}):
            async with httpx.AsyncClient(transport=httpx.ASGITransport(app=create_app(self.runner)),base_url='http://127.0.0.1:8000') as c:
                r=await c.post('/api/backgrounds',content=b'x',headers={'Origin':'http://localhost:5173','Content-Length':str(REQUEST_LIMIT+1)})
                self.assertEqual(r.status_code,413)
                self.assertEqual(r.headers['access-control-allow-origin'],'http://localhost:5173')

if __name__=='__main__': unittest.main(verbosity=2)

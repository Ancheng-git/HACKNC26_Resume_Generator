#!/usr/bin/env python3
"""Fixed ATS layout using the same content JSON as build_professional.py.

Usage: python scripts/build_ats_resume.py professional_Company.json output/Resume.docx
       python scripts/build_ats_resume.py professional_Company.json --verify ATS.docx Professional.docx
Optional ATS-only fields: technical_skills/additional (string arrays), education
dates/details, and experience kind (employment/training/project/competition).
Existing Professional JSON files remain valid. No model calls are made here.
"""
import argparse
from collections import Counter
import json
from pathlib import Path
import re
import tempfile
import unicodedata

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

try:
    from .build_professional import validate
    from .check_pages import check
except ImportError:
    from build_professional import validate
    from check_pages import check


def plain(text):
    return text.replace("**", "")


def _construct(c):
    # Retain the approved plain Figma layout: US Letter, Calibri 10.5 pt,
    # centered identity, black section rules and right-aligned dates.
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    s.left_margin = s.right_margin = Inches(0.6)
    s.top_margin = s.bottom_margin = Inches(0.5)
    normal = d.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    def para(text="", bold=False, size=10.5, before=0, after=1.5, center=False, bullet=False):
        p = d.add_paragraph()
        pf = p.paragraph_format
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        pf.widow_control = True
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if bullet:
            pf.left_indent, pf.first_line_indent = Inches(0.2), Inches(-0.15)
            p.add_run("•  ")
        for i, part in enumerate(re.split(r"\*\*", text)):
            if not part:
                continue
            r = p.add_run(part)
            r.font.name, r.font.size = "Calibri", Pt(size)
            r._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Calibri")
            r.bold = bold or i % 2 == 1
        return p

    def heading(text):
        p = para(text, bold=True, size=11, before=6, after=2)
        p.paragraph_format.keep_with_next = True
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for key, value in (("val", "single"), ("sz", "4"), ("space", "1"), ("color", "000000")):
            bottom.set(qn("w:" + key), value)
        border.append(bottom)
        p._p.get_or_add_pPr().append(border)

    def dated(text, dates):
        p = para(text)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.3), WD_TAB_ALIGNMENT.RIGHT)
        if dates:
            p.add_run("\t" + dates)

    para(c["name"], bold=True, size=15, after=0, center=True)
    para(" | ".join(c["contact"]), center=True)
    para(c["headline"], bold=True, size=11, after=0, center=True)
    heading("PROFILE")
    for text in c["profile"]:
        para(text)
    heading("CORE COMPETENCIES")
    para(" | ".join(c["core_expertise"]))
    heading("EDUCATION")
    for item in c["education"]:
        dated("**" + item["school"] + "**", item.get("dates", ""))
        para(item["degree"])
        for text in item.get("details", []):
            para(text)
    # Keep source order; add a truthful category heading only when it changes.
    category = None
    for job in c["experience"]:
        kind = job.get("kind", "employment")
        current = "EXPERIENCE" if kind == "employment" else "TRAINING, PROJECTS, AND COMPETITIONS"
        if current != category:
            heading(current)
            category = current
        dated("**" + job["title"] + "** | " + job["company"], job["dates"])
        for text in job["bullets"]:
            para(text, bullet=True)
    skills = c.get("technical_skills", [])
    certificates = c.get("certifications", [])
    additional = c.get("additional", [])
    if skills or certificates or additional:
        heading("TECHNICAL SKILLS")
        for text in skills:
            para(text)
        if certificates:
            para("**Certifications:** " + "; ".join(certificates))
        for text in additional:
            para(text)
    return d


def build(content, out_path):
    validate(content, require_photo=False)
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="ats-build-", dir=out_path.parent) as tmp:
        candidate = Path(tmp) / out_path.name
        _construct(content).save(candidate)
        check(candidate)
        candidate.replace(out_path)
    return out_path


def verify_shared_content(content, paths):
    """Check selected identity, role fields, profile and bullet text survive both layouts.

    This proves consistency with the selected JSON, not truth of the source facts.
    """
    validate(content, require_photo=False)

    def normalized(text):
        return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", plain(text))).strip().lstrip("•▪ ")

    expected = Counter(normalized(s) for s in content["profile"] + [
        b for job in content["experience"] for b in job["bullets"]])
    metadata = [content["name"], content["headline"]] + [
        job[field] for job in content["experience"] for field in ("company", "title", "dates") if job[field]]
    warnings = []
    for path in paths:
        doc = Document(path)
        paragraphs = [normalized(p.text) for p in doc.paragraphs]
        actual = Counter(paragraphs)
        if expected - actual:
            raise ValueError("Selected profile or accomplishment bullets are missing or changed")
        text = "\n".join(paragraphs)
        if any(normalized(value).casefold() not in text.casefold() for value in metadata):
            raise ValueError("Selected identity, headline, role title or dates are missing or changed")
    # The first path is ATS. Short content is a review warning, never a padding mandate.
    text = " ".join(p.text for p in Document(paths[0]).paragraphs)
    count = len(re.findall(r"\S+", text))
    if count < 350:
        warnings.append(f"The ATS resume has about {count} words, below the suggested length. Review any omitted relevant evidence. Keep it brief if the source material is limited; do not invent or repeat content to fill space.")
    return warnings


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("content", type=Path)
    parser.add_argument("output", nargs="?", type=Path)
    parser.add_argument("--verify", nargs="+", type=Path)
    args = parser.parse_args()
    content = json.loads(args.content.read_text(encoding="utf-8"))
    if args.verify:
        if args.output:
            parser.error("output and --verify cannot be combined")
        print(json.dumps({"warnings": verify_shared_content(content, args.verify)}, ensure_ascii=False))
    elif args.output:
        print(f"Saved {build(content, args.output)} (1 page)")
    else:
        parser.error("supply an output path or --verify document paths")


if __name__ == "__main__":
    main()

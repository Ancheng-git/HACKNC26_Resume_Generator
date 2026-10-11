"""Regression checks with synthetic content; never invoke Claude or a model API."""
import asyncio
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import subprocess
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from PIL import Image

from scripts import ats_check, build_ats_resume, build_professional
from scripts import build_cover_letter_professional as letter
from backend.runner import Runner, TaskError

ROOT = Path(__file__).resolve().parent.parent


def content():
    return {
        'name': 'Synthetic Alex', 'contact': ['alex@example.test'],
        'headline': 'ANALYST CANDIDATE',
        'profile': ['Synthetic analyst using Python to explain financial datasets.'],
        'core_expertise': ['Python', 'Data analysis'],
        'education': [{'degree': 'B.S. Statistics, in progress', 'school': 'Fixture University',
                       'dates': 'Expected 2029', 'details': ['Coursework: Statistics']}],
        'experience': [
            {'company': 'Fixture Team', 'title': 'Analyst Intern', 'dates': '2025–2026',
             'bullets': ['Analyzed **42 records** with Python and explained findings.']},
            {'company': 'Fixture Program', 'title': 'Training Fellow', 'dates': '2026', 'kind': 'training',
             'bullets': ['Completed **8** synthetic training exercises.']}],
        'technical_skills': ['Languages: Python, SQL'],
        'certifications': ['Synthetic course certificate'],
    }


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.base = Path(self.tmp.name).resolve()
        self.photo = self.base / 'photo.jpg'
        Image.new('RGB', (100, 120), 'gray').save(self.photo)

    def tearDown(self):
        self.tmp.cleanup()

    def test_keywords_reject_empty_and_malformed_but_allow_no_preferences(self):
        for value in ({}, {'must_have': [], 'nice_to_have': []},
                      {'must_have': 'Python'}, {'must_have': [' ']}):
            with self.assertRaises(ValueError):
                ats_check.check_keywords('Python', value)
        result = ats_check.check_keywords('Python', {'must_have': ['Python'], 'nice_to_have': []})
        self.assertEqual(result['must_have_pct'], 100)

    def test_sections_require_actual_heading_lines(self):
        missing, _ = ats_check.check_sections('My education and experience developed core competencies and a profile.')
        self.assertEqual(len(missing), 4)
        self.assertEqual(ats_check.check_sections('PROFILE\nText\nEXPERIENCE:\nText\nEducation\nText\nTechnical Skills')[0], [])

    def test_nested_types_fail_before_rendering_and_old_shape_stays_valid(self):
        c = content()
        build_professional.validate(c, require_photo=False)
        old = deepcopy(c)
        old.pop('technical_skills')
        old['education'][0].pop('details')
        old['education'][0].pop('dates')
        old['experience'][1].pop('kind')
        build_professional.validate(old, require_photo=False)
        for field, value in [('name', ' '), ('profile', 'not a list'), ('experience', 'not a list'),
                             ('education', [{'degree': 'B.S.'}])]:
            bad = deepcopy(c); bad[field] = value
            with self.assertRaises(ValueError):
                build_professional.validate(bad, require_photo=False)
        bad = deepcopy(c); bad['experience'][0]['bullets'] = [' ']
        with self.assertRaises(ValueError):
            build_professional.validate(bad, require_photo=False)

    def test_js_leading_bold_and_middle_bold_match_markup(self):
        # Execute the actual helper with a tiny TextRun stand-in, not a second implementation.
        source = (ROOT / 'scripts/build_cover_letter.js').read_text()
        helper = source[source.index('function richRuns('):source.index('function sidebarHeading(')]
        program = 'class TextRun { constructor(o) { Object.assign(this,o); } }\n' + helper + '''
console.log(JSON.stringify(['**50+** posts','Analyzed **42** records','plain'].map(text=>richRuns(text,{size:21,color:'000000'}))));
'''
        runs = json.loads(subprocess.check_output(['node', '-e', program], text=True))
        self.assertEqual([(r['text'], r['bold']) for r in runs[0]], [('50+', True), (' posts', False)])
        self.assertEqual([r['bold'] for r in runs[1]], [False, True, False])
        self.assertFalse(runs[2][0]['bold'])

    def test_real_three_builders_preserve_shared_content_and_one_page(self):
        c = content()
        paths = [self.base / n for n in ('ATS.docx', 'Professional.docx', 'Letter.docx')]
        with patch.object(build_professional, 'PHOTO', self.photo), patch.object(letter, 'PHOTO', self.photo), \
             patch.dict(letter.validate.__globals__, {'PHOTO': self.photo}):
            build_ats_resume.build(c, paths[0])
            build_professional.build(c, paths[1])
            letter.build({**c, 'date': 'October 10, 2026', 'salutation': 'Dear Team,',
                          'paragraphs': ['**42 projects** demonstrate synthetic analysis experience.']}, paths[2])
        self.assertTrue(build_ats_resume.verify_shared_content(c, paths[:2]))
        doc = Document(paths[0])
        text = '\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('TRAINING, PROJECTS, AND COMPETITIONS', text)
        self.assertIn('Expected 2029', text)
        self.assertIn('Languages: Python, SQL', text)
        self.assertEqual(ats_check.check_format(paths[0], doc)[0], [])
        # A metric alteration or an omitted selected bullet must fail even if ATS keywords match.
        paragraph = next(p for p in doc.paragraphs if '42 records' in p.text)
        paragraph.text = paragraph.text.replace('42 records', '43 records')
        doc.save(paths[0])
        with self.assertRaises(ValueError):
            build_ats_resume.verify_shared_content(c, paths[:2])

    def test_overflow_does_not_overwrite_existing_output(self):
        c = content(); target = self.base / 'ATS.docx'
        build_ats_resume.build(c, target)
        before = hashlib.sha256(target.read_bytes()).hexdigest()
        c['experience'][0]['bullets'] *= 100
        with self.assertRaises(RuntimeError):
            build_ats_resume.build(c, target)
        self.assertEqual(hashlib.sha256(target.read_bytes()).hexdigest(), before)

    def test_ats_publication_requires_both_thresholds_or_completed_gap_review(self):
        directory = self.base
        (directory / 'output').mkdir(); (directory / 'preview').mkdir()
        names = dict(zip(('ats_resume', 'professional_resume', 'professional_cover_letter'),
                         ('ATS.docx', 'Professional.docx', 'Letter.docx')))
        reviews = {}
        for kind, name in names.items():
            path = directory / 'output' / name; path.write_bytes(b'isolated ATS validation fixture')
            image = 'preview/' + kind + '.png'; (directory / image).write_bytes(b'\x89PNG\r\n\x1a\n')
            reviews[kind] = {'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                             'image': image, 'notes': 'Synthetic review fixture; no model'}
        (directory / 'report.md').write_text('Synthetic pass/gap report fixture')
        (directory / 'keywords_Test.json').write_text('{}')
        (directory / 'professional_Test.json').write_text('{}')
        runner = Runner(root=ROOT, runs=directory)

        async def verify():
            for composite, coverage, passes, gaps, accepted in (
                    (93, 83, 0, 'Unsupported required tool', False),
                    (93, 83, 2, 'Unsupported required tool', False),
                    (93, 83, 3, 'Unsupported required tool after three rechecks', True),
                    (93, 83, 3, '', False),
                    (79, 100, 0, '', False),
                    (80, 90, 0, '', True)):
                with self.subTest(composite=composite, coverage=coverage, passes=passes, gaps=gaps):
                    (directory / 'completion.json').write_text(json.dumps({
                        'visual_review': reviews, 'revision_passes': passes, 'gaps': gaps}))
                    async def fake_process(args, cwd, label, input_text=None):
                        if label == 'ats':
                            (cwd / 'ats.stdout.log').write_text(json.dumps({
                                'composite_score': composite, 'keyword_match': {'must_have_pct': coverage},
                                'format': {'issues': []}}))
                        if label == 'content':
                            (cwd / 'content.stdout.log').write_text(json.dumps({'warnings': []}))
                        return 0
                    task = {}
                    with patch.object(runner, '_process', side_effect=fake_process):
                        if accepted:
                            await runner._validate(task, directory, names, 'keywords_Test.json')
                            self.assertEqual(len(task['files']), 3)
                            self.assertEqual(bool(task['warnings']), coverage < 90)
                        else:
                            with self.assertRaisesRegex(TaskError, rf'composite {composite} .*coverage {coverage}%') as caught:
                                await runner._validate(task, directory, names, 'keywords_Test.json')
                            self.assertIn(f'{passes}/3', str(caught.exception))
                            self.assertNotIn('files', task)
            await runner.close()
        asyncio.run(verify())

    def test_format_violation_cannot_be_overridden_by_perfect_score(self):
        directory = self.base
        names = dict(zip(('ats_resume', 'professional_resume', 'professional_cover_letter'),
                         ('ATS.docx', 'Professional.docx', 'Letter.docx')))
        (directory / 'output').mkdir(); (directory / 'preview').mkdir()
        reviews = {}
        for kind, name in names.items():
            path = directory / 'output' / name; path.write_bytes(b'isolated validation fixture')
            image = 'preview/' + kind + '.png'; (directory / image).write_bytes(b'\x89PNG\r\n\x1a\n')
            reviews[kind] = {'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
                             'image': image, 'notes': 'Simulated review; no model'}
        (directory / 'completion.json').write_text(json.dumps({'visual_review': reviews, 'revision_passes': 3, 'gaps': 'fixture'}))
        (directory / 'report.md').write_text('Synthetic fixture')
        (directory / 'keywords_Test.json').write_text('{}')
        runner = Runner(root=ROOT, runs=directory)

        async def fake_process(args, cwd, label, input_text=None):
            if label == 'ats':
                (cwd / 'ats.stdout.log').write_text(json.dumps({
                    'composite_score': 100, 'keyword_match': {'must_have_pct': 100},
                    'format': {'issues': ['Contains a forbidden table']}}))
            return 0

        async def verify():
            with patch.object(runner, '_process', side_effect=fake_process):
                with self.assertRaisesRegex(TaskError, 'format violations'):
                    await runner._validate({}, directory, names, 'keywords_Test.json')
            await runner.close()
        asyncio.run(verify())


if __name__ == '__main__':
    unittest.main()
